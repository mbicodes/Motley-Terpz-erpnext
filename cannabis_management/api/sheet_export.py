import datetime
from io import BytesIO

import frappe
import openpyxl
from frappe.desk.utils import provide_binary_file
from frappe.utils.csvutils import to_csv
from frappe.utils.xlsxutils import ILLEGAL_CHARACTERS_RE, get_excel_date_format, handle_html
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.workbook.child import INVALID_TITLE_REGEX

MAX_SHEET_NAME_LENGTH = 31


def add_sheet(wb, rows, sheet_name, column_widths=None):
	sheet_name = INVALID_TITLE_REGEX.sub(" ", sheet_name)[:MAX_SHEET_NAME_LENGTH]
	ws = wb.create_sheet(sheet_name)
	ws.row_dimensions[1].font = Font(name="Calibri", bold=True)

	for i, width in enumerate(column_widths or [], start=1):
		if width:
			ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

	date_format, time_format = get_excel_date_format()

	for row in rows:
		clean_row = []
		for item in row:
			if isinstance(item, str):
				value = handle_html(item)
				if next(ILLEGAL_CHARACTERS_RE.finditer(value), None):
					value = ILLEGAL_CHARACTERS_RE.sub("", value)
			else:
				value = item

			if isinstance(value, datetime.date | datetime.datetime):
				number_format = date_format
				if isinstance(value, datetime.datetime):
					number_format = f"{date_format} {time_format}"
				cell = WriteOnlyCell(ws, value=value)
				cell.number_format = number_format
				clean_row.append(cell)
			else:
				clean_row.append(value)

		ws.append(clean_row)


def send_xlsx(title, sheets):
	wb = openpyxl.Workbook(write_only=True)
	written = 0
	for sheet in sheets:
		name, rows = sheet[0], sheet[1]
		widths = sheet[2] if len(sheet) > 2 else None
		if len(rows) < 2:
			continue
		add_sheet(wb, rows, name, widths)
		written += 1

	if not written:
		frappe.throw(frappe._("There is nothing to export."))

	out = BytesIO()
	wb.save(out)
	provide_binary_file(title, "xlsx", out.getvalue())


def send_csv(title, rows):
	if len(rows) < 2:
		frappe.throw(frappe._("There is nothing to export."))

	clean = []
	for row in rows:
		clean.append([
			"" if cell is None
			else cell.strftime("%Y-%m-%d") if isinstance(cell, datetime.date | datetime.datetime)
			else cell
			for cell in row
		])

	provide_binary_file(title, "csv", to_csv(clean).encode("utf-8-sig"))
